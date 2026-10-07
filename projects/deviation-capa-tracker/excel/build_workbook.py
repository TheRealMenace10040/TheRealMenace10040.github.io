"""Build Deviation_CAPA_Tracker.xlsx from deviations.json.

deviations.json is the same seeded dataset the web tracker generates (exported from index.html),
so the workbook and the web page show the same records. Every calculated column and dashboard
figure is written as an Excel formula, not a value.

Usage: python build_workbook.py deviations.json Deviation_CAPA_Tracker.xlsx
"""
import json, sys, datetime as dt
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

D=json.load(open(sys.argv[1])); OUT=sys.argv[2]
P=lambda s: dt.date.fromisoformat(s) if s else None
recs=D['recs']; ASOF=P(D['asof'])
LAST=400  # formula rows pre-filled to this row so new deviations can be added
F0=5

INK='121518'; MUTED='596166'; HAIR='CFD4D0'; TINT='F3F5F3'; PEN='1D3D73'
F=lambda **k: Font(name='Arial',size=k.pop('size',10),**k)
HDR_FILL=PatternFill('solid',fgColor=INK); TINT_FILL=PatternFill('solid',fgColor=TINT)
INPUT_FILL=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color=HAIR); rule=Side(style='medium',color=INK)
BOX=Border(bottom=thin)
BLUE=F(color='0000FF'); GREEN=F(color='008000')

wb=Workbook()

def title(ws,t,sub):
    ws['A1']=t; ws['A1'].font=F(size=16,bold=True)
    ws['A2']=sub; ws['A2'].font=F(size=9,color=MUTED,italic=True)
    ws.sheet_view.showGridLines=False
def header(ws,row,col,labels,widths=None):
    for i,h in enumerate(labels):
        c=ws.cell(row=row,column=col+i,value=h)
        c.font=F(bold=True,color='FFFFFF',size=9); c.fill=HDR_FILL
        c.alignment=Alignment(wrap_text=True,vertical='center')
        if widths: ws.column_dimensions[L(col+i)].width=widths[i]
    ws.row_dimensions[row].height=30

# ---------------- Instructions ----------------
ws=wb.active; ws.title='Instructions'
title(ws,'Deviation and CAPA Tracker','QA-DEV-006 · Simulated data for a portfolio demonstration. No real site data.')
ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=95
rows=[
 ('What this is','An Excel deviation and CAPA tracker. Every status, age, repeat flag, CAPA effectiveness result and dashboard number is a formula, so it updates when you add records, close them, or change the data date.'),
 ('',''),
 ('SHEETS',''),
 ('Settings','Data date, closure target and the other rules. Change these and the whole workbook follows. Also holds the drop-down lists.'),
 ('Deviation Log','One row per deviation. Enter data in the blue-text columns A to J; columns K to V are calculated. Formulas are filled down to row %d.'%LAST),
 ('Dashboard','KPIs, aging by severity, monthly close-out trend with chart, root-cause Pareto and a summary by area.'),
 ('Repeat Analysis','Equipment / root-cause pairs that have repeated, with record count, repeats, last occurrence and CAPA effectiveness.'),
 ('',''),
 ('HOW TO USE',''),
 ('Add a deviation','Type in the next empty row of Deviation Log: ID, opened date, area, equipment, severity, root cause and owner. Area, severity and root cause have drop-downs.'),
 ('Close a deviation','Enter the closed date in column H. Status, days open and overdue flag update automatically.'),
 ('Raise a CAPA','Enter the CAPA number in column I. CAPA due date = deviation closed + the CAPA due days on Settings.'),
 ('Close a CAPA','Enter the CAPA closed date in column J. The effectiveness result starts as "Check pending" and becomes "Effective" after the check window, or "Ineffective" if the same root cause comes back on the same equipment.'),
 ('Roll forward','Change "Data as of" on Settings to today (or use =TODAY()).'),
 ('',''),
 ('COLOUR KEY',''),
 ('Blue text','Inputs you type in.'),
 ('Black text','Formulas. Do not overwrite.'),
 ('Yellow fill','Settings you can change.'),
 ('Red row','Open and past the closure target (overdue).'),
 ('Amber row','Open and past the warning threshold (due soon).'),
 ('',''),
 ('DEFINITIONS',''),
 ('Overdue','Open, and days open is greater than the closure target.'),
 ('Repeat','Same root cause on the same equipment / process within the repeat window before this deviation was opened.'),
 ('CAPA effectiveness','Ineffective if the same root cause recurs on the same equipment / process within the effectiveness window after the CAPA closed. Effective once the window has passed with no recurrence. Otherwise Check pending.'),
 ('Median days to close','Median of days from opened to closed, for records closed in the period. Median is used because a few long investigations would distort an average.'),
]
for i,(a,b) in enumerate(rows,start=4):
    ws.cell(row=i,column=1,value=a).font=F(bold=a.isupper() or bool(b),size=10 if not a.isupper() else 9,color=MUTED if a.isupper() else INK)
    c=ws.cell(row=i,column=2,value=b); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(row=i,column=1).alignment=Alignment(vertical='top')
ws['A20'].font=F(bold=True,color='0000FF'); ws['A22'].fill=INPUT_FILL
ws['A23'].fill=PatternFill('solid',fgColor='F6DEDB'); ws['A24'].fill=PatternFill('solid',fgColor='F6ECD2')
for r in range(4,4+len(rows)):
    b=ws.cell(row=r,column=2).value or ''
    ws.row_dimensions[r].height=max(15,15*(len(b)//95+1))

# ---------------- Settings ----------------
st=wb.create_sheet('Settings')
title(st,'Settings','Change the yellow cells. Everything else in the workbook reads from here.')
st.column_dimensions['A'].width=4; st.column_dimensions['B'].width=38; st.column_dimensions['C'].width=14
settings=[('Data as of','AsOf',ASOF,'dd-mmm-yyyy','Date the tracker is reported as of. Use =TODAY() for a live tracker.'),
 ('Closure target (days)','Target',30,'0','Deviations should be investigated and closed within this many calendar days.'),
 ('Due-soon warning (days)','Warn',20,'0','Open records past this age are highlighted amber.'),
 ('CAPA due after deviation close (days)','CapaDue',90,'0','CAPA due date = deviation closed date + this.'),
 ('Repeat window (days)','RepeatWin',180,'0','Same cause on the same equipment within this window counts as a repeat.'),
 ('Effectiveness check window (days)','EffWin',180,'0','CAPA is judged effective if the cause does not recur within this window after CAPA closure.')]
header(st,4,2,['Setting','Value'])
for i,(lab,name,val,fmt,note) in enumerate(settings,start=5):
    st.cell(row=i,column=2,value=lab).font=F()
    c=st.cell(row=i,column=3,value=val); c.font=BLUE; c.fill=INPUT_FILL; c.number_format=fmt; c.border=BOX
    c.comment=Comment(note,'Tracker')
    st.cell(row=i,column=4,value=note).font=F(size=9,color=MUTED)
    wb.defined_names[name]=DefinedName(name,attr_text=f"Settings!$C${i}")
st.column_dimensions['D'].width=70

lists={'Areas':list(D['areas'].keys()),'Severities':['Minor','Major','Critical'],
       'RootCauses':sorted({c for v in D['causes'].values() for c in v}),
       'Owners':sorted({r['owner'] for r in recs})}
lr=13
st.cell(row=lr-1,column=2,value='Drop-down lists (add new values to the bottom of a list)').font=F(bold=True,size=9,color=MUTED)
for j,(name,vals) in enumerate(lists.items()):
    c=2+j
    st.column_dimensions[L(c)].width=max(st.column_dimensions[L(c)].width or 10,30 if j else 38)
    h=st.cell(row=lr,column=c,value=name); h.font=F(bold=True,color='FFFFFF',size=9); h.fill=HDR_FILL
    for k,v in enumerate(vals): st.cell(row=lr+1+k,column=c,value=v).font=BLUE
    wb.defined_names['List'+name]=DefinedName('List'+name,attr_text=f"Settings!${L(c)}${lr+1}:${L(c)}${lr+30}")

# ---------------- Deviation Log ----------------
lg=wb.create_sheet('Deviation Log')
title(lg,'Deviation Log','Blue columns A to J are inputs. Grey-headed columns K to V are calculated; do not type over them.')
cols=[('Deviation ID',13),('Opened',11),('Area',16),('Equipment / process',24),('Severity',9),('Root cause',27),('Owner',11),('Closed',11),('CAPA ID',10),('CAPA closed',11),
      ('Status',8),('Days open',8),('Overdue',8),('Aging bucket (open)',9),('Repeat',7),('CAPA due',11),('CAPA status',13),('CAPA effectiveness',13),('Opened month',9),('Closed month',9),('Days to close, last 90 d',10),('Days to close, prior 6 mo',10)]
header(lg,4,1,[c for c,_ in cols],[w for _,w in cols])
for i in range(11,23): lg.cell(row=4,column=i).fill=PatternFill('solid',fgColor='596166')
R=lambda c: f"${c}${F0}:${c}${LAST}"
for i in range(F0,LAST+1):
    r=recs[i-F0] if i-F0<len(recs) else None
    if r:
        vals=[r['id'],P(r['opened']),r['area'],r['asset'],r['sev'],r['cause'],r['owner'],P(r['closed']),r['capa'],P(r['capaClosed'])]
        for j,v in enumerate(vals,start=1):
            c=lg.cell(row=i,column=j,value=v); c.font=BLUE
    f={
     'K':f'=IF(A{i}="","",IF(H{i}="","Open","Closed"))',
     'L':f'=IF(A{i}="","",IF(H{i}="",AsOf-B{i},H{i}-B{i}))',
     'M':f'=IF(A{i}="","",IF(AND(K{i}="Open",L{i}>Target),"Yes","No"))',
     'N':f'=IF(K{i}<>"Open","",IF(L{i}<=30,"0-30",IF(L{i}<=60,"31-60",IF(L{i}<=90,"61-90","90+"))))',
     'O':f'=IF(A{i}="","",IF(COUNTIFS({R("D")},D{i},{R("F")},F{i},{R("B")},">="&(B{i}-RepeatWin),{R("B")},"<"&B{i})>0,"Yes","No"))',
     'P':f'=IF(OR(I{i}="",H{i}=""),"",H{i}+CapaDue)',
     'Q':f'=IF(A{i}="","",IF(I{i}="","Not required",IF(J{i}<>"","Closed",IF(H{i}="","Awaiting dev. close",IF(AsOf>P{i},"Overdue","Open")))))',
     'R':f'=IF(J{i}="","",IF(COUNTIFS({R("D")},D{i},{R("F")},F{i},{R("B")},">"&J{i},{R("B")},"<="&(J{i}+EffWin))>0,"Ineffective",IF(AsOf-J{i}>=EffWin,"Effective","Check pending")))',
     'S':f'=IF(A{i}="","",DATE(YEAR(B{i}),MONTH(B{i}),1))',
     'T':f'=IF(H{i}="","",DATE(YEAR(H{i}),MONTH(H{i}),1))',
     'U':f'=IF(H{i}="","",IF(AsOf-H{i}<=90,L{i},""))',
     'V':f'=IF(H{i}="","",IF(AND(AsOf-H{i}>90,AsOf-H{i}<=270),L{i},""))'}
    for k,v in f.items(): lg[f'{k}{i}']=v; lg[f'{k}{i}'].font=F()
    for k in 'BHJP': lg[f'{k}{i}'].number_format='dd-mmm-yy'
    for k in 'ST': lg[f'{k}{i}'].number_format='mmm-yy'
lg.freeze_panes='B5'
lg.auto_filter.ref=f"A4:V{F0+len(recs)-1}"
for col_,lst in (('C','ListAreas'),('E','ListSeverities'),('F','ListRootCauses'),('G','ListOwners')):
    dv=DataValidation(type='list',formula1='='+lst,allow_blank=True,showErrorMessage=(col_!='G'),errorTitle='Not in list',error='Pick a value from the list, or add it on the Settings sheet first.')
    lg.add_data_validation(dv); dv.add(f'{col_}{F0}:{col_}{LAST}')
dvd=DataValidation(type='date',operator='greaterThan',formula1='36526',allow_blank=True,showErrorMessage=True,error='Enter a date.')
lg.add_data_validation(dvd); dvd.add(f'B{F0}:B{LAST}'); dvd.add(f'H{F0}:H{LAST}'); dvd.add(f'J{F0}:J{LAST}')
rng=f"A{F0}:V{LAST}"
lg.conditional_formatting.add(rng,FormulaRule(formula=[f'$M{F0}="Yes"'],fill=PatternFill('solid',fgColor='F6DEDB'),stopIfTrue=True))
lg.conditional_formatting.add(rng,FormulaRule(formula=[f'AND($K{F0}="Open",$L{F0}>Warn)'],fill=PatternFill('solid',fgColor='F6ECD2'),stopIfTrue=True))
lg.conditional_formatting.add(f'O{F0}:O{LAST}',CellIsRule(operator='equal',formula=['"Yes"'],font=Font(name='Arial',bold=True,color='9A6A00')))
lg.conditional_formatting.add(f'R{F0}:R{LAST}',CellIsRule(operator='equal',formula=['"Ineffective"'],font=Font(name='Arial',bold=True,color='B02A20')))
lg.conditional_formatting.add(f'Q{F0}:Q{LAST}',CellIsRule(operator='equal',formula=['"Overdue"'],font=Font(name='Arial',bold=True,color='B02A20')))
lg.conditional_formatting.add(f'M{F0}:M{LAST}',CellIsRule(operator='equal',formula=['"Yes"'],font=Font(name='Arial',bold=True,color='B02A20')))

# ---------------- Dashboard ----------------
db=wb.create_sheet('Dashboard',1)
title(db,'Deviation and CAPA Dashboard','All figures are formulas on the Deviation Log and Settings sheets.')
db['A3']='=" Data as of "&TEXT(AsOf,"dd mmm yyyy")&"  ·  closure target "&Target&" days"'; db['A3'].font=F(size=9,color=MUTED)
for c,w in zip('ABCDEFGHIJ',[30,12,12,12,12,12,4,30,12,12]): db.column_dimensions[c].width=w
LG="'Deviation Log'!"
LR=lambda c: f"{LG}${c}${F0}:${c}${LAST}"
kpis=[('Total deviations',f'=COUNTA({LR("A")})','0'),
 ('Open deviations',f'=COUNTIF({LR("K")},"Open")','0'),
 ('Overdue (open past target)',f'=COUNTIF({LR("M")},"Yes")','0'),
 ('Critical deviations open',f'=COUNTIFS({LR("E")},"Critical",{LR("K")},"Open")','0'),
 ('Median days to close, last 90 days',f'=IFERROR(MEDIAN({LR("U")}),"-")','0.0'),
 ('Median days to close, prior 6 months',f'=IFERROR(MEDIAN({LR("V")}),"-")','0.0'),
 ('Repeat deviations',f'=COUNTIF({LR("O")},"Yes")','0'),
 ('Repeat rate',f'=IFERROR(B11/B5,0)','0.0%'),
 ('CAPAs raised',f'=COUNTIF({LR("I")},"?*")','0'),
 ('CAPAs overdue',f'=COUNTIF({LR("Q")},"Overdue")','0'),
 ('CAPAs with effectiveness result',f'=COUNTIF({LR("R")},"Effective")+COUNTIF({LR("R")},"Ineffective")','0'),
 ('CAPAs ineffective',f'=COUNTIF({LR("R")},"Ineffective")','0'),
 ('CAPA effectiveness rate',f'=IFERROR((B15-B16)/B15,0)','0.0%')]
header(db,4,1,['Key figures','Value'])
for i,(a,f,fmt) in enumerate(kpis,start=5):
    db.cell(row=i,column=1,value=a).font=F(); c=db.cell(row=i,column=2,value=f); c.font=F(bold=True); c.number_format=fmt
    for cc in (1,2): db.cell(row=i,column=cc).border=BOX
db.conditional_formatting.add('B7',CellIsRule(operator='greaterThan',formula=['0'],font=Font(name='Arial',bold=True,color='B02A20')))
db.conditional_formatting.add('B9',CellIsRule(operator='greaterThan',formula=['Target'],font=Font(name='Arial',bold=True,color='B02A20')))
db.conditional_formatting.add('B16',CellIsRule(operator='greaterThan',formula=['0'],font=Font(name='Arial',bold=True,color='B02A20')))

# aging grid
db['H4']='Aging of open deviations (days open)'; db['H4'].font=F(bold=True)
header(db,5,8,['Severity','0-30','31-60','61-90','90+','Total'])
for c in 'KLM': db.column_dimensions[c].width=10
for j,s in enumerate(['Critical','Major','Minor','All'],start=6):
    db.cell(row=j,column=8,value=s).font=F(bold=True)
    for k,b in enumerate(['0-30','31-60','61-90','90+'],start=9):
        sev='' if s=='All' else f'{LR("E")},"{s}",'
        db.cell(row=j,column=k,value=f'=COUNTIFS({sev}{LR("N")},"{b}")').font=F(bold=(s=='All'))
    db.cell(row=j,column=13,value=f'=SUM(I{j}:L{j})').font=F(bold=True)
    for k in range(8,14): db.cell(row=j,column=k).border=BOX; db.cell(row=j,column=k).alignment=Alignment(horizontal='center') if k>8 else Alignment()
db.conditional_formatting.add('J6:L8',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='F6DEDB'),font=Font(name='Arial',bold=True,color='B02A20')))

# area summary
db['H12']='By area'; db['H12'].font=F(bold=True)
header(db,13,8,['Area','Total','Open','Overdue','Repeats','Critical'])
for j,a in enumerate(lists['Areas'],start=14):
    db.cell(row=j,column=8,value=a).font=F()
    for k,expr in enumerate([f'=COUNTIF({LR("C")},H{j})',f'=COUNTIFS({LR("C")},H{j},{LR("K")},"Open")',f'=COUNTIFS({LR("C")},H{j},{LR("M")},"Yes")',f'=COUNTIFS({LR("C")},H{j},{LR("O")},"Yes")',f'=COUNTIFS({LR("C")},H{j},{LR("E")},"Critical")'],start=9):
        c=db.cell(row=j,column=k,value=expr); c.font=F(); c.alignment=Alignment(horizontal='center')
    for k in range(8,14): db.cell(row=j,column=k).border=BOX

# monthly trend
tr=21
db.cell(row=tr,column=1,value='Monthly close-out trend').font=F(bold=True)
header(db,tr+1,1,['Month','Opened','Closed','Median days to close','Target (days)','Open at month end'])
months=[]; y,m=2025,4
while (y,m)<=(ASOF.year,ASOF.month):
    months.append(dt.date(y,m,1)); m+=1
    if m>12: y,m=y+1,1
for j,mo in enumerate(months,start=tr+2):
    c=db.cell(row=j,column=1,value=mo); c.number_format='mmm-yy'; c.font=BLUE; c.alignment=Alignment(horizontal='left')
    db.cell(row=j,column=2,value=f'=COUNTIF({LR("S")},A{j})')
    db.cell(row=j,column=3,value=f'=COUNTIF({LR("T")},A{j})')
    db.cell(row=j,column=4).value=ArrayFormula(f'D{j}',f'=IFERROR(MEDIAN(IF({LR("T")}=A{j},{LR("L")})),NA())')
    db.cell(row=j,column=5,value='=Target')
    db.cell(row=j,column=6,value=f'=COUNTIFS({LR("B")},"<="&EOMONTH(A{j},0),{LR("A")},"<>")-COUNTIFS({LR("H")},"<="&EOMONTH(A{j},0))')
    for k in range(1,7): db.cell(row=j,column=k).border=BOX; db.cell(row=j,column=k).font=F() if k>1 else BLUE
    db.cell(row=j,column=4).number_format='0.0'
mend=tr+1+len(months)
db.cell(row=mend+1,column=1,value='Months are inputs: extend the list down and copy the formulas to add a month. If nothing closed in a month the median is left blank as a chart gap (NA).').font=F(size=8,color=MUTED,italic=True)

bar=BarChart(); bar.type='col'; bar.grouping='clustered'; bar.title='Opened vs closed per month, with median days to close'
bar.y_axis.title='Deviations'; bar.height=9; bar.width=22
bar.add_data(Reference(db,min_col=2,min_row=tr+1,max_col=3,max_row=mend),titles_from_data=True)
bar.set_categories(Reference(db,min_col=1,min_row=tr+2,max_row=mend))
bar.series[0].graphicalProperties.solidFill='9AA1A6'; bar.series[1].graphicalProperties.solidFill=PEN
bar.x_axis.number_format='mmm-yy'
ln=LineChart(); ln.add_data(Reference(db,min_col=4,min_row=tr+1,max_col=5,max_row=mend),titles_from_data=True)
ln.y_axis.axId=200; ln.y_axis.title='Median days to close'; ln.y_axis.crosses='max'
ln.series[0].graphicalProperties.line.solidFill='B02A20'; ln.series[0].graphicalProperties.line.width=28000
ln.series[1].graphicalProperties.line.solidFill='B02A20'; ln.series[1].graphicalProperties.line.dashStyle='dash'; ln.series[1].graphicalProperties.line.width=12000
bar.x_axis.delete=False; bar.y_axis.delete=False; ln.y_axis.delete=False
for se in ln.series: se.smooth=False
bar.y_axis.scaling.min=0; bar.y_axis.scaling.max=16; bar.y_axis.majorUnit=2
ln.y_axis.scaling.min=0; ln.y_axis.scaling.max=50; ln.y_axis.majorUnit=10
bar.y_axis.majorGridlines=None
bar+=ln; bar.legend.position='b'
db.add_chart(bar,f'H{tr}')

# pareto
pr=mend+4
db.cell(row=pr,column=1,value='Root cause Pareto').font=F(bold=True)
header(db,pr+1,1,['Root cause','Count','% of total','Cumulative %'])
order=[c for c,_ in Counter(r['cause'] for r in recs).most_common()]
for c in lists['RootCauses']:
    if c not in order: order.append(c)
for j,c in enumerate(order,start=pr+2):
    db.cell(row=j,column=1,value=c).font=F()
    db.cell(row=j,column=2,value=f'=COUNTIF({LR("F")},A{j})').font=F()
    x=db.cell(row=j,column=3,value=f'=IFERROR(B{j}/$B$5,0)'); x.number_format='0.0%'; x.font=F()
    y_=db.cell(row=j,column=4,value=f'=SUM($C${pr+2}:C{j})'); y_.number_format='0.0%'; y_.font=F()
    for k in range(1,5): db.cell(row=j,column=k).border=BOX
pend=pr+1+len(order)
db.cell(row=pend+1,column=1,value='Causes are listed in order of count for the sample data. Re-sort (Data > Sort by Count) after adding records.').font=F(size=8,color=MUTED,italic=True)
pb=BarChart(); pb.type='bar'; pb.title='Deviations by root cause'; pb.height=12; pb.width=22; pb.legend=None
pb.add_data(Reference(db,min_col=2,min_row=pr+1,max_row=pend),titles_from_data=True)
pb.set_categories(Reference(db,min_col=1,min_row=pr+2,max_row=pend))
pb.series[0].graphicalProperties.solidFill=INK; pb.x_axis.scaling.orientation='maxMin'
pb.x_axis.delete=False; pb.y_axis.delete=False; pb.x_axis.tickLblSkip=1; pb.gapWidth=60
db.add_chart(pb,f'H{pr}')
db.freeze_panes='A4'

# ---------------- Repeat Analysis ----------------
ra=wb.create_sheet('Repeat Analysis')
title(ra,'Repeat Analysis','Equipment / root-cause pairs that have repeated. Counts and CAPA results are formulas on the Deviation Log.')
hd=['Area','Equipment / process','Root cause','Records','Repeats','First opened','Last opened','CAPAs raised','CAPAs ineffective','CAPAs open','Flag']
header(ra,4,1,hd,[16,26,28,9,9,12,12,10,11,10,20])
groups=defaultdict(list)
for r in recs: groups[(r['area'],r['asset'],r['cause'])].append(r)
# which pairs to list: same rules as the Repeat and CAPA effectiveness columns, run in Python only to choose and order rows
def has_repeat(v):
    d=sorted(P(x['opened']) for x in v)
    return any(0<=(b-a).days<=180 and a!=b for a,b in zip(d,d[1:]))
def ineffective(v):
    return sum(1 for x in v if x['capaClosed'] and any(0<(P(y['opened'])-P(x['capaClosed'])).days<=180 for y in v))
keys=[k for k,v in groups.items() if has_repeat(v)]
keys.sort(key=lambda k:(-ineffective(groups[k]),-len(groups[k]),k))
crit=lambda j: f'{LR("D")},B{j},{LR("F")},C{j}'
for j,(a,asset,cause) in enumerate(keys,start=5):
    for k,v in enumerate([a,asset,cause],start=1): ra.cell(row=j,column=k,value=v).font=BLUE
    ra.cell(row=j,column=4,value=f'=COUNTIFS({crit(j)})')
    ra.cell(row=j,column=5,value=f'=COUNTIFS({crit(j)},{LR("O")},"Yes")')
    ra.cell(row=j,column=6,value=f'=_xlfn.MINIFS({LR("B")},{crit(j)})').number_format='dd-mmm-yy'
    ra.cell(row=j,column=7,value=f'=_xlfn.MAXIFS({LR("B")},{crit(j)})').number_format='dd-mmm-yy'
    ra.cell(row=j,column=8,value=f'=COUNTIFS({crit(j)},{LR("I")},"?*")')
    ra.cell(row=j,column=9,value=f'=COUNTIFS({crit(j)},{LR("R")},"Ineffective")')
    ra.cell(row=j,column=10,value=f'=COUNTIFS({crit(j)},{LR("I")},"?*",{LR("Q")},"<>Closed")')
    ra.cell(row=j,column=11,value=f'=IF(I{j}>0,"CAPA ineffective",IF(J{j}>0,"CAPA open",IF(H{j}=0,"No CAPA yet","Check pending")))')
    for k in range(4,12): ra.cell(row=j,column=k).font=F(); ra.cell(row=j,column=k).border=BOX
    for k in range(1,4): ra.cell(row=j,column=k).border=BOX
rend=4+len(keys)
ra.conditional_formatting.add(f'A5:K{rend}',FormulaRule(formula=['$I5>0'],fill=PatternFill('solid',fgColor='F6DEDB')))
ra.conditional_formatting.add(f'K5:K{rend}',CellIsRule(operator='equal',formula=['"CAPA ineffective"'],font=Font(name='Arial',bold=True,color='B02A20')))
ra.cell(row=rend+2,column=1,value='Pairs are listed for the sample data. To check a new pair, type its area, equipment and root cause in the next row and copy the formulas down.').font=F(size=8,color=MUTED,italic=True)
ra.freeze_panes='A5'

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.font and c.font.name!='Arial': c.font=Font(name='Arial',size=c.font.size or 10,bold=c.font.bold,italic=c.font.italic,color=c.font.color)
    s.sheet_properties.pageSetUpPr=None
    s.page_setup.orientation='landscape'; s.page_setup.fitToWidth=1
wb.save(OUT); print('saved',OUT,len(recs),'records',len(keys),'repeat pairs')
