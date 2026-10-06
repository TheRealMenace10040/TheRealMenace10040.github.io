"""Project 2: Environmental Monitoring QC Trend Report (synthetic data).
Generates data -> loads SQLite -> runs SQL summaries -> builds formula-driven Excel report."""
import os, csv, random, sqlite3, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

random.seed(7)
OUT = "."
os.makedirs(f"{OUT}/data", exist_ok=True)
os.makedirs(f"{OUT}/sql", exist_ok=True)

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="375623")
RED = PatternFill("solid", fgColor="F8CBAD")
YEL = PatternFill("solid", fgColor="FFE699")

# sample type, alert limit, action limit (CFU) - illustrative values only
LIMITS = [("Active air", 5, 10), ("Surface contact", 3, 5), ("Compressed air", 1, 3), ("Personnel glove", 1, 3)]
ROOMS = [("CR-101", "ISO 7"), ("CR-102", "ISO 7"), ("CR-201", "ISO 8"), ("CR-202", "ISO 8"), ("Gowning", "ISO 8")]
ROOM_BIAS = {"CR-101": 0.6, "CR-102": 0.8, "CR-201": 1.0, "CR-202": 1.5, "Gowning": 1.3}

rows = []
start = dt.date(2026, 1, 5)
sid = 1
for d in range(0, 266):
    day = start + dt.timedelta(days=d)
    if day.weekday() >= 5:
        continue
    for room, iso in ROOMS:
        for stype, alert, action in LIMITS:
            if stype == "Compressed air" and room == "Gowning":
                continue
            lam = {"Active air": 1.3, "Surface contact": 0.55, "Compressed air": 0.04, "Personnel glove": 0.06}[stype] * ROOM_BIAS[room]
            # seasonal bump in summer months for realism
            if day.month in (7, 8):
                lam *= 1.4
            cfu = sum(1 for _ in range(30) if random.random() < lam / 30)
            rows.append((f"EM{sid:06d}", day.isoformat(), room, iso, stype, cfu,
                         random.choice(["Tech A", "Tech B", "Tech C", "Tech D"])))
            sid += 1

with open(f"{OUT}/data/em_samples.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["sample_id", "sample_date", "room", "iso_class", "sample_type", "cfu", "analyst"])
    w.writerows(rows)
with open(f"{OUT}/data/limits.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["sample_type", "alert_limit", "action_limit"])
    w.writerows(LIMITS)

# ---------------- SQL -----------------
SCHEMA = """-- Environmental monitoring schema (SQLite / PostgreSQL compatible)
CREATE TABLE limits (
    sample_type   TEXT PRIMARY KEY,
    alert_limit   INTEGER NOT NULL,
    action_limit  INTEGER NOT NULL
);
CREATE TABLE em_samples (
    sample_id    TEXT PRIMARY KEY,
    sample_date  DATE NOT NULL,
    room         TEXT NOT NULL,
    iso_class    TEXT NOT NULL,
    sample_type  TEXT NOT NULL REFERENCES limits(sample_type),
    cfu          INTEGER NOT NULL CHECK (cfu >= 0),
    analyst      TEXT
);
CREATE INDEX ix_em_date ON em_samples(sample_date);
CREATE INDEX ix_em_room ON em_samples(room, sample_type);
"""
QUERIES = """-- 1. Classify every result against its limits
CREATE VIEW v_results AS
SELECT s.*,
       CASE WHEN s.cfu >= l.action_limit THEN 'Action'
            WHEN s.cfu >= l.alert_limit  THEN 'Alert'
            ELSE 'Pass' END AS result_status
FROM em_samples s
JOIN limits l USING (sample_type);

-- 2. Monthly excursion rate by room
-- name: monthly_room_excursions
SELECT strftime('%Y-%m', sample_date)                          AS month,
       room,
       COUNT(*)                                                AS samples,
       SUM(result_status = 'Alert')                            AS alerts,
       SUM(result_status = 'Action')                           AS actions,
       ROUND(100.0 * SUM(result_status <> 'Pass') / COUNT(*), 2) AS excursion_pct
FROM v_results
GROUP BY month, room
ORDER BY month, room;

-- 3. Rooms whose excursion rate in the last 30 days is above their prior 90-day baseline
-- name: rooms_trending_up
WITH recent AS (
    SELECT room, AVG(result_status <> 'Pass') AS rate
    FROM v_results
    WHERE sample_date >  date((SELECT MAX(sample_date) FROM em_samples), '-30 day')
    GROUP BY room
), baseline AS (
    SELECT room, AVG(result_status <> 'Pass') AS rate
    FROM v_results
    WHERE sample_date <= date((SELECT MAX(sample_date) FROM em_samples), '-30 day')
      AND sample_date >  date((SELECT MAX(sample_date) FROM em_samples), '-120 day')
    GROUP BY room
)
SELECT r.room,
       ROUND(100 * b.rate, 2) AS baseline_pct,
       ROUND(100 * r.rate, 2) AS last30_pct,
       ROUND(100 * (r.rate - b.rate), 2) AS change_pts
FROM recent r JOIN baseline b USING (room)
WHERE r.rate > b.rate
ORDER BY change_pts DESC;

-- 4. Repeat excursions: same room + sample type exceeding a limit 3+ times in 14 days
-- name: repeat_excursions
SELECT a.room, a.sample_type, a.sample_date,
       COUNT(*) AS excursions_in_14_days
FROM v_results a
JOIN v_results b
  ON  b.room = a.room AND b.sample_type = a.sample_type
  AND b.result_status <> 'Pass'
  AND b.sample_date BETWEEN date(a.sample_date, '-13 day') AND a.sample_date
WHERE a.result_status <> 'Pass'
GROUP BY a.sample_id
HAVING COUNT(*) >= 3
ORDER BY a.sample_date;
"""
open(f"{OUT}/sql/schema.sql", "w").write(SCHEMA)
open(f"{OUT}/sql/analysis_queries.sql", "w").write(QUERIES)

db = sqlite3.connect(":memory:")
db.executescript(SCHEMA)
db.executemany("INSERT INTO limits VALUES (?,?,?)", LIMITS)
db.executemany("INSERT INTO em_samples VALUES (?,?,?,?,?,?,?)", rows)
blocks = QUERIES.split("\n\n-- ")
db.executescript(blocks[0])  # view
results = {}
for b in blocks[1:]:
    name = [l for l in b.splitlines() if l.startswith("-- name:")][0].split(":")[1].strip()
    sql = "\n".join(l for l in b.splitlines() if not l.strip().startswith("--") and not l[:1].isdigit())
    cur = db.execute(sql)
    results[name] = ([c[0] for c in cur.description], cur.fetchall())
for name, (cols, data) in results.items():
    with open(f"{OUT}/data/{name}.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(cols); w.writerows(data)
    print(name, len(data), data[:3])

# ---------------- Excel -----------------
wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)


def header(ws, cols, widths, fill=FILL):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HDR, fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


rd = wb.active
rd.title = "README"
for line in [
    "Environmental Monitoring QC Trend Report",
    "Portfolio project by Dennis Nguyen. ALL DATA IS SYNTHETIC; limits are illustrative, not from any real SOP.",
    "",
    "Pipeline: Python generates ~4,000 EM results -> SQLite (sql/schema.sql) -> SQL analysis (sql/analysis_queries.sql)",
    "-> this workbook, where formulas classify every result and build the monthly summary and charts.",
    "",
    "Sheets:",
    "  RawData        - every sample; formulas look up alert/action limits and flag Pass / Alert / Action.",
    "  Limits         - alert and action limits by sample type (edit here, everything recalculates).",
    "  MonthlySummary - samples, alerts, actions, excursion % per room per month (COUNTIFS).",
    "  RoomTrend      - SQL output: rooms whose last-30-day excursion rate is above their 90-day baseline.",
    "  RepeatExcursions - SQL output: same room + sample type over a limit 3+ times in 14 days.",
    "  Weekly         - filled by the BuildWeeklyReport macro.",
    "",
    "Automation: import vba/EMReport.bas, save as .xlsm, then run BuildWeeklyReport. It asks for a week-ending date,",
    "copies that week's results to the Weekly sheet, totals excursions by room, highlights Action results, and saves a PDF.",
]:
    rd.append([line])
rd["A1"].font = Font(bold=True, size=14)
rd.column_dimensions["A"].width = 115

lim = wb.create_sheet("Limits")
header(lim, ["SampleType", "AlertLimit", "ActionLimit"], [18, 11, 12])
for l in LIMITS:
    lim.append(list(l))

raw = wb.create_sheet("RawData")
header(raw, ["SampleID", "Date", "Room", "ISO", "SampleType", "CFU", "Analyst", "AlertLimit", "ActionLimit", "Result", "Month"],
       [11, 11, 9, 7, 16, 6, 9, 10, 11, 9, 10])
for i, r in enumerate(rows, start=2):
    raw.append([r[0], dt.date.fromisoformat(r[1]), r[2], r[3], r[4], r[5], r[6]])
    raw[f"H{i}"] = f"=INDEX(Limits!$B$2:$B$5,MATCH(E{i},Limits!$A$2:$A$5,0))"
    raw[f"I{i}"] = f"=INDEX(Limits!$C$2:$C$5,MATCH(E{i},Limits!$A$2:$A$5,0))"
    raw[f"J{i}"] = f'=IF(F{i}>=I{i},"Action",IF(F{i}>=H{i},"Alert","Pass"))'
    raw[f"K{i}"] = f"=DATE(YEAR(B{i}),MONTH(B{i}),1)"
    raw[f"B{i}"].number_format = "yyyy-mm-dd"
    raw[f"K{i}"].number_format = "mmm yyyy"
N = len(rows) + 1
t = Table(displayName="tblEM", ref=f"A1:K{N}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight11", showRowStripes=True)
raw.add_table(t)
raw.conditional_formatting.add(f"J2:J{N}", CellIsRule(operator="equal", formula=['"Action"'], fill=RED))
raw.conditional_formatting.add(f"J2:J{N}", CellIsRule(operator="equal", formula=['"Alert"'], fill=YEL))

ms = wb.create_sheet("MonthlySummary", 1)
months = sorted({dt.date.fromisoformat(r[1]).replace(day=1) for r in rows})
ms["A1"] = "Excursion rate (% of samples at Alert or Action) by room and month"
ms["A1"].font = Font(bold=True, size=13)
ms.append([])
ms.append(["Month"] + [r for r, _ in ROOMS] + ["All rooms", "Action results"])
for c in ms[3]:
    c.font, c.fill = HDR, FILL
for j, m in enumerate(months):
    rr = 4 + j
    ms[f"A{rr}"] = m
    ms[f"A{rr}"].number_format = "mmm yyyy"
    for k, (room, _) in enumerate(ROOMS):
        col = chr(ord("B") + k)
        ms[f"{col}{rr}"] = (f'=IFERROR(COUNTIFS(RawData!$K$2:$K${N},$A{rr},RawData!$C$2:$C${N},{col}$3,RawData!$J$2:$J${N},"<>Pass")'
                            f'/COUNTIFS(RawData!$K$2:$K${N},$A{rr},RawData!$C$2:$C${N},{col}$3),"")')
        ms[f"{col}{rr}"].number_format = "0.0%"
    ms[f"G{rr}"] = (f'=IFERROR(COUNTIFS(RawData!$K$2:$K${N},$A{rr},RawData!$J$2:$J${N},"<>Pass")'
                    f'/COUNTIFS(RawData!$K$2:$K${N},$A{rr}),"")')
    ms[f"G{rr}"].number_format = "0.0%"
    ms[f"H{rr}"] = f'=COUNTIFS(RawData!$K$2:$K${N},$A{rr},RawData!$J$2:$J${N},"Action")'
last = 3 + len(months)
for col in "ABCDEFGH":
    ms.column_dimensions[col].width = 12
ms.conditional_formatting.add(f"B4:G{last}", CellIsRule(operator="greaterThan", formula=["0.1"], fill=RED))

lc = LineChart()
lc.title = "Monthly excursion rate by room"
lc.y_axis.number_format = "0%"
lc.add_data(Reference(ms, min_col=2, max_col=6, min_row=3, max_row=last), titles_from_data=True)
lc.set_categories(Reference(ms, min_col=1, min_row=4, max_row=last))
lc.height, lc.width = 8, 18
ms.add_chart(lc, f"A{last + 3}")
bc = BarChart()
bc.title = "Action-level results per month"
bc.add_data(Reference(ms, min_col=8, min_row=3, max_row=last), titles_from_data=True)
bc.set_categories(Reference(ms, min_col=1, min_row=4, max_row=last))
bc.height, bc.width = 8, 14
ms.add_chart(bc, f"J{last + 3}")

for name, title in (("rooms_trending_up", "RoomTrend"), ("repeat_excursions", "RepeatExcursions")):
    ws = wb.create_sheet(title)
    cols, data = results[name]
    header(ws, cols, [16] * len(cols))
    for r in data:
        ws.append(list(r))

wk = wb.create_sheet("Weekly")
wk["A1"] = "Run BuildWeeklyReport (vba/EMReport.bas) to fill this sheet."

wb.save(f"{OUT}/EM_QC_Trend_Report.xlsx")
print("rows", len(rows), "months", len(months))
