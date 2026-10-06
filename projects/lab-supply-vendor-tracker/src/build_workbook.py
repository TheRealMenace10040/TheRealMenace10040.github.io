"""Project 1: Lab Supply & Vendor Performance Tracker (synthetic data)."""
import os, random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

random.seed(42)
OUT = "."
os.makedirs(OUT, exist_ok=True)
TODAY = dt.date(2026, 10, 1)  # fixed "as of" date so results are reproducible

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
RED = PatternFill("solid", fgColor="F8CBAD")
YEL = PatternFill("solid", fgColor="FFE699")
GRN = PatternFill("solid", fgColor="C6E0B4")

vendors = [
    ("V01", "Northstar Media Supply", "Culture media", 10, 0.06),
    ("V02", "Pacific BioReagents", "Reagents", 14, 0.18),
    ("V03", "CleanLab Consumables", "Consumables", 7, 0.05),
    ("V04", "Sterile Source Inc.", "Sterility supplies", 12, 0.10),
    ("V05", "Coastal Gas & Air", "Compressed gas", 5, 0.25),
    ("V06", "Apex Calibration Svcs", "Calibration", 21, 0.30),
    ("V07", "BioIndicator Labs", "Biological indicators", 18, 0.12),
    ("V08", "Summit Glassware", "Glassware", 9, 0.08),
]
# sku, desc, vendor, unit cost, shelf life days (0 = no expiry), min stock, reorder qty
items = [
    ("MED-001", "Tryptic Soy Agar plates (10/pk)", "V01", 28.50, 90, 40, 60),
    ("MED-002", "Sabouraud Dextrose Agar plates (10/pk)", "V01", 31.00, 90, 20, 30),
    ("MED-003", "R2A Agar plates (10/pk)", "V01", 34.25, 60, 15, 25),
    ("MED-004", "Contact plates, TSA w/ neutralizer (20/pk)", "V01", 46.00, 75, 30, 40),
    ("MED-005", "Fluid Thioglycollate tubes (25/pk)", "V01", 52.75, 120, 10, 20),
    ("RGT-001", "Gram stain kit", "V02", 89.00, 365, 4, 6),
    ("RGT-002", "Endotoxin LAL reagent kit", "V02", 412.00, 180, 3, 5),
    ("RGT-003", "Catalase reagent (100 mL)", "V02", 24.50, 270, 5, 10),
    ("RGT-004", "ID panel cartridges (20/box)", "V02", 265.00, 240, 6, 8),
    ("CON-001", "Sterile pipette tips 1000 uL (960/cs)", "V03", 118.00, 0, 8, 12),
    ("CON-002", "Sterile swabs (100/pk)", "V03", 19.75, 730, 15, 25),
    ("CON-003", "Nitrile gloves M (1000/cs)", "V03", 72.00, 0, 10, 15),
    ("CON-004", "Membrane filters 0.45 um (100/pk)", "V03", 96.50, 1095, 6, 10),
    ("STR-001", "Sterile sampling bags (500/cs)", "V04", 145.00, 730, 4, 6),
    ("STR-002", "Isopropyl alcohol 70% sterile (12/cs)", "V04", 88.00, 540, 6, 10),
    ("STR-003", "Sterile gowns (30/cs)", "V04", 132.00, 1095, 5, 8),
    ("GAS-001", "Nitrogen cylinder, UHP", "V05", 64.00, 0, 3, 4),
    ("GAS-002", "CO2 cylinder, incubator grade", "V05", 58.00, 0, 3, 4),
    ("CAL-001", "Pipette calibration service (per unit)", "V06", 45.00, 0, 0, 20),
    ("CAL-002", "Air sampler calibration service", "V06", 380.00, 0, 0, 2),
    ("BIO-001", "Steam BI ampoules (100/box)", "V07", 310.00, 365, 3, 4),
    ("BIO-002", "EO BI strips (100/box)", "V07", 275.00, 365, 2, 3),
    ("GLS-001", "Media bottles 500 mL (10/cs)", "V08", 64.00, 0, 4, 6),
    ("GLS-002", "Petri dish glass (12/cs)", "V08", 58.50, 0, 3, 5),
]
vendor_lt = {v[0]: (v[3], v[4]) for v in vendors}

wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HDR, FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


# --- README sheet --------------------------------------------------------
rd = wb.active
rd.title = "README"
for line in [
    "Lab Supply & Vendor Performance Tracker",
    "Portfolio project by Dennis Nguyen. ALL DATA IS SYNTHETIC (randomly generated); vendor names are fictional.",
    "",
    "Problem: a QC microbiology lab depends on time-sensitive materials (media with 60-120 day shelf life,",
    "biological indicators, reagents). Late deliveries and expiring lots cause testing delays and waste.",
    "",
    "What this workbook does:",
    "  PurchaseOrders  - 12 months of POs; formulas compute lead time, on-time flag, and days late.",
    "  Inventory       - lots on hand; formulas compute days to expiry and Expired / Expiring / OK status.",
    "  ReorderAlerts   - usable (non-expired) stock vs. minimum; flags items to reorder and suggests qty.",
    "  VendorScorecard - orders, spend, avg lead time, on-time % per vendor (COUNTIFS/SUMIFS/AVERAGEIFS).",
    "  Dashboard       - headline KPIs, on-time % by vendor chart, monthly spend trend.",
    "",
    "Automation: import vba/VendorTracker.bas (Alt+F11 > File > Import) and save as .xlsm. Macros:",
    "  RefreshAll_And_Stamp - recalculates, sorts the reorder list, and timestamps the dashboard.",
    "  ExportReorderList    - writes the items needing reorder to a dated CSV for purchasing.",
    "  ImportPOs            - appends new PO rows from a CSV export into the PurchaseOrders table.",
    "",
    f"'As of' date for all calculations is fixed at {TODAY.isoformat()} (cell Dashboard!B2) so results are reproducible.",
]:
    rd.append([line])
rd["A1"].font = Font(bold=True, size=14)
rd.column_dimensions["A"].width = 110

# --- Vendors ---------------------------------------------------------------
ws = wb.create_sheet("Vendors")
header(ws, ["VendorID", "VendorName", "Category", "QuotedLeadDays"], [10, 26, 20, 14])
for v in vendors:
    ws.append(list(v[:4]))
table(ws, "tblVendors", f"A1:D{len(vendors)+1}")

# --- Items -----------------------------------------------------------------
ws = wb.create_sheet("Items")
header(ws, ["SKU", "Description", "VendorID", "UnitCost", "ShelfLifeDays", "MinStock", "ReorderQty"],
       [10, 42, 10, 10, 13, 10, 11])
for it in items:
    ws.append(list(it))
for r in range(2, len(items) + 2):
    ws.cell(row=r, column=4).number_format = "$#,##0.00"
table(ws, "tblItems", f"A1:G{len(items)+1}")
NI = len(items) + 1

# --- PurchaseOrders -----------------------------------------------------------
ws = wb.create_sheet("PurchaseOrders")
cols = ["PO", "SKU", "VendorID", "OrderDate", "PromisedDate", "ReceivedDate", "Qty",
        "UnitCost", "LineTotal", "LeadTimeDays", "OnTime", "DaysLate", "OrderMonth"]
header(ws, cols, [10, 10, 10, 12, 13, 13, 6, 10, 11, 12, 9, 9, 11])
start = dt.date(2025, 10, 1)
rows = []
for n in range(320):
    sku = random.choice(items)
    vid = sku[2]
    quoted, late_p = vendor_lt[vid]
    od = start + dt.timedelta(days=random.randint(0, 360))
    prom = od + dt.timedelta(days=quoted)
    if random.random() < late_p:
        actual = quoted + random.randint(1, 9)
    else:
        actual = quoted - random.randint(0, min(3, quoted - 1))
    rec = od + dt.timedelta(days=actual)
    rows.append((od, sku[0], prom, rec if rec <= TODAY else None, random.randint(1, 6) * max(1, sku[6] // 4)))
rows.sort()
for i, (od, s, prom, rec, q) in enumerate(rows, start=2):
    ws.append([f"PO-{24000+i}", s, None, od, prom, rec, q])
    ws[f"C{i}"] = f'=INDEX(Items!$C$2:$C${NI},MATCH(B{i},Items!$A$2:$A${NI},0))'
    ws[f"H{i}"] = f'=INDEX(Items!$D$2:$D${NI},MATCH(B{i},Items!$A$2:$A${NI},0))'
    ws[f"I{i}"] = f"=G{i}*H{i}"
    ws[f"J{i}"] = f'=IF(F{i}="","",F{i}-D{i})'
    ws[f"K{i}"] = f'=IF(F{i}="","Open",IF(F{i}<=E{i},"Yes","No"))'
    ws[f"L{i}"] = f'=IF(F{i}="","",MAX(0,F{i}-E{i}))'
    ws[f"M{i}"] = f'=DATE(YEAR(D{i}),MONTH(D{i}),1)'
    for c in "DEF":
        ws[f"{c}{i}"].number_format = "yyyy-mm-dd"
    ws[f"M{i}"].number_format = "mmm yyyy"
    ws[f"H{i}"].number_format = ws[f"I{i}"].number_format = "$#,##0.00"
NP = len(rows) + 1
table(ws, "tblPOs", f"A1:M{NP}")
ws.conditional_formatting.add(f"K2:K{NP}", CellIsRule(operator="equal", formula=['"No"'], fill=RED))

# --- Inventory ----------------------------------------------------------------
ws = wb.create_sheet("Inventory")
header(ws, ["Lot", "SKU", "QtyOnHand", "ReceivedDate", "ExpiryDate", "DaysToExpiry", "Status"],
       [12, 10, 11, 13, 12, 13, 11])
r = 2
for it in items:
    for _ in range(random.randint(1, 3)):
        recd = TODAY - dt.timedelta(days=random.randint(5, 150))
        ws.append([f"L{random.randint(100000, 999999)}", it[0], random.randint(0, it[6]), recd])
        ws[f"E{r}"] = f'=IF(INDEX(Items!$E$2:$E${NI},MATCH(B{r},Items!$A$2:$A${NI},0))=0,"",D{r}+INDEX(Items!$E$2:$E${NI},MATCH(B{r},Items!$A$2:$A${NI},0)))'
        ws[f"F{r}"] = f'=IF(E{r}="","",E{r}-Dashboard!$B$2)'
        ws[f"G{r}"] = f'=IF(E{r}="","No expiry",IF(F{r}<0,"Expired",IF(F{r}<=30,"Expiring","OK")))'
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "yyyy-mm-dd"
        r += 1
NV = r - 1
table(ws, "tblInventory", f"A1:G{NV}")
ws.conditional_formatting.add(f"G2:G{NV}", CellIsRule(operator="equal", formula=['"Expired"'], fill=RED))
ws.conditional_formatting.add(f"G2:G{NV}", CellIsRule(operator="equal", formula=['"Expiring"'], fill=YEL))

# --- ReorderAlerts ----------------------------------------------------------------
ws = wb.create_sheet("ReorderAlerts")
header(ws, ["SKU", "Description", "VendorID", "UsableQty", "MinStock", "OpenPOQty", "Reorder?", "SuggestedQty"],
       [10, 42, 10, 11, 10, 11, 10, 13])
for i, it in enumerate(items, start=2):
    ws[f"A{i}"] = it[0]
    ws[f"B{i}"] = f"=INDEX(Items!$B$2:$B${NI},MATCH(A{i},Items!$A$2:$A${NI},0))"
    ws[f"C{i}"] = f"=INDEX(Items!$C$2:$C${NI},MATCH(A{i},Items!$A$2:$A${NI},0))"
    # usable = not expired
    ws[f"D{i}"] = f'=SUMIFS(Inventory!$C$2:$C${NV},Inventory!$B$2:$B${NV},A{i},Inventory!$G$2:$G${NV},"<>Expired")'
    ws[f"E{i}"] = f"=INDEX(Items!$F$2:$F${NI},MATCH(A{i},Items!$A$2:$A${NI},0))"
    ws[f"F{i}"] = f'=SUMIFS(PurchaseOrders!$G$2:$G${NP},PurchaseOrders!$B$2:$B${NP},A{i},PurchaseOrders!$K$2:$K${NP},"Open")'
    ws[f"G{i}"] = f'=IF(D{i}+F{i}<E{i},"YES","no")'
    ws[f"H{i}"] = f'=IF(G{i}="YES",MAX(INDEX(Items!$G$2:$G${NI},MATCH(A{i},Items!$A$2:$A${NI},0)),E{i}-D{i}-F{i}),0)'
NR = len(items) + 1
table(ws, "tblReorder", f"A1:H{NR}")
ws.conditional_formatting.add(f"A2:H{NR}", FormulaRule(formula=[f'$G2="YES"'], fill=RED))

# --- VendorScorecard ------------------------------------------------------------
ws = wb.create_sheet("VendorScorecard")
header(ws, ["VendorID", "VendorName", "Orders", "Received", "OnTime", "OnTime%", "AvgLeadDays",
            "QuotedLeadDays", "AvgDaysLate", "Spend12mo", "Rating"],
       [10, 26, 9, 10, 9, 10, 12, 14, 12, 13, 14])
P = "PurchaseOrders!"
for i, v in enumerate(vendors, start=2):
    ws[f"A{i}"] = v[0]
    ws[f"B{i}"] = f"=INDEX(Vendors!$B$2:$B$9,MATCH(A{i},Vendors!$A$2:$A$9,0))"
    ws[f"C{i}"] = f"=COUNTIFS({P}$C$2:$C${NP},A{i})"
    ws[f"D{i}"] = f'=COUNTIFS({P}$C$2:$C${NP},A{i},{P}$K$2:$K${NP},"<>Open")'
    ws[f"E{i}"] = f'=COUNTIFS({P}$C$2:$C${NP},A{i},{P}$K$2:$K${NP},"Yes")'
    ws[f"F{i}"] = f'=IF(D{i}=0,"",E{i}/D{i})'
    ws[f"G{i}"] = f'=IFERROR(AVERAGEIFS({P}$J$2:$J${NP},{P}$C$2:$C${NP},A{i},{P}$K$2:$K${NP},"<>Open"),"")'
    ws[f"H{i}"] = f"=INDEX(Vendors!$D$2:$D$9,MATCH(A{i},Vendors!$A$2:$A$9,0))"
    ws[f"I{i}"] = f'=IFERROR(AVERAGEIFS({P}$L$2:$L${NP},{P}$C$2:$C${NP},A{i},{P}$K$2:$K${NP},"<>Open"),"")'
    ws[f"J{i}"] = f"=SUMIFS({P}$I$2:$I${NP},{P}$C$2:$C${NP},A{i})"
    ws[f"K{i}"] = f'=IF(F{i}="","",IF(F{i}>=0.9,"Preferred",IF(F{i}>=0.8,"Monitor","At risk")))'
    ws[f"F{i}"].number_format = "0.0%"
    ws[f"G{i}"].number_format = ws[f"I{i}"].number_format = "0.0"
    ws[f"J{i}"].number_format = "$#,##0"
table(ws, "tblScorecard", "A1:K9")
ws.conditional_formatting.add("K2:K9", CellIsRule(operator="equal", formula=['"At risk"'], fill=RED))
ws.conditional_formatting.add("K2:K9", CellIsRule(operator="equal", formula=['"Monitor"'], fill=YEL))
ws.conditional_formatting.add("K2:K9", CellIsRule(operator="equal", formula=['"Preferred"'], fill=GRN))

# --- Dashboard -------------------------------------------------------------------
ws = wb.create_sheet("Dashboard", 1)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 16
ws["A1"] = "Lab Supply & Vendor Dashboard (synthetic data)"
ws["A1"].font = Font(bold=True, size=14)
kpis = [
    ("As-of date", TODAY, "yyyy-mm-dd"),
    ("Total spend (12 mo)", f"=SUM({P}I2:I{NP})", "$#,##0"),
    ("Purchase orders", f"=COUNTA({P}A2:A{NP})", "0"),
    ("Overall on-time delivery %", f'=COUNTIF({P}K2:K{NP},"Yes")/COUNTIF({P}K2:K{NP},"<>Open")', "0.0%"),
    ("Avg lead time (days)", f'=AVERAGEIF({P}K2:K{NP},"<>Open",{P}J2:J{NP})', "0.0"),
    ("Open POs", f'=COUNTIF({P}K2:K{NP},"Open")', "0"),
    ("Items below reorder point", f'=COUNTIF(ReorderAlerts!G2:G{NR},"YES")', "0"),
    ("Lots expired", f'=COUNTIF(Inventory!G2:G{NV},"Expired")', "0"),
    ("Lots expiring within 30 days", f'=COUNTIF(Inventory!G2:G{NV},"Expiring")', "0"),
    ("Last refreshed (macro)", "", "@"),
]
for i, (label, val, fmt) in enumerate(kpis, start=2):
    ws[f"A{i}"] = label
    ws[f"A{i}"].font = Font(bold=True)
    ws[f"B{i}"] = val
    ws[f"B{i}"].number_format = fmt

# monthly spend helper table
ws["D1"] = "Month"; ws["E1"] = "Spend"
for c in ("D1", "E1"):
    ws[c].font, ws[c].fill = HDR, FILL
for m in range(12):
    r = m + 2
    first = dt.date(2025 + (9 + m) // 12, (9 + m) % 12 + 1, 1)
    ws[f"D{r}"] = first
    ws[f"D{r}"].number_format = "mmm yy"
    ws[f"E{r}"] = f"=SUMIFS({P}$I$2:$I${NP},{P}$M$2:$M${NP},D{r})"
    ws[f"E{r}"].number_format = "$#,##0"
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 12

bar = BarChart()
bar.title = "On-time delivery % by vendor"
bar.y_axis.number_format = "0%"
bar.add_data(Reference(wb["VendorScorecard"], min_col=6, min_row=1, max_row=9), titles_from_data=True)
bar.set_categories(Reference(wb["VendorScorecard"], min_col=1, min_row=2, max_row=9))
bar.height, bar.width = 7.5, 16
ws.add_chart(bar, "A14")

line = LineChart()
line.title = "Monthly spend"
line.add_data(Reference(ws, min_col=5, min_row=1, max_row=13), titles_from_data=True)
line.set_categories(Reference(ws, min_col=4, min_row=2, max_row=13))
line.height, line.width = 7.5, 16
ws.add_chart(line, "G14")

wb.save(f"{OUT}/Lab_Supply_Vendor_Tracker.xlsx")

# CSV sample for ImportPOs macro
with open(f"{OUT}/sample_new_POs.csv", "w") as f:
    f.write("PO,SKU,OrderDate,PromisedDate,ReceivedDate,Qty\n")
    f.write("PO-25001,MED-001,2026-09-28,2026-10-08,,30\n")
    f.write("PO-25002,BIO-001,2026-09-29,2026-10-17,,2\n")
    f.write("PO-25003,GAS-002,2026-09-30,2026-10-05,,4\n")
print("ok", NP, NV)
